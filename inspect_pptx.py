import pptx

prs = pptx.Presentation(r'C:\Users\tk706\.gemini\antigravity\brain\b36f7999-cc10-492d-88b9-6d654c5f7704\.user_uploaded\media_1792306310983_19193e02.pptx')
slide1 = prs.slides[0]
for shape in slide1.shapes:
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            print(f'P: text="{p.text}"')
            for r in p.runs:
                c = r.font.color.rgb if r.font.color and r.font.color.type == 1 else None
                print(f'  Run: "{r.text}", font={r.font.name}, size={r.font.size}, bold={r.font.bold}, color={c}')

slide3 = prs.slides[2]
print("\n=== Slide 3 Title ===")
for shape in slide3.shapes:
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            print(f'P: text="{p.text}"')
            for r in p.runs:
                c = r.font.color.rgb if r.font.color and r.font.color.type == 1 else None
                print(f'  Run: "{r.text}", font={r.font.name}, size={r.font.size}, bold={r.font.bold}, color={c}')
